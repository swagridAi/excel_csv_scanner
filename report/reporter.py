# report/reporter.py
"""
Reporter module for generating analysis reports in various formats.

This module provides functionality to generate reports from file analysis results
in different formats including Excel, CSV, and HTML.
"""
import os
import logging
import pandas as pd
import json
from datetime import datetime
from pathlib import Path
from typing import List, Dict, Any, Optional, Union, Callable, Set

# Try to import openpyxl for Excel report formatting
try:
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
    from openpyxl.utils import get_column_letter
    OPENPYXL_AVAILABLE = True
except ImportError:
    OPENPYXL_AVAILABLE = False

from config import REPORT_COLUMN_ORDER, EXCEL_HEADER_STYLE

# Configure logger
logger = logging.getLogger(__name__)


class Reporter:
    """
    Generate reports from file analysis results.
    
    This class handles the creation of reports in various formats,
    organizing and formatting the data for better readability.
    """
    
    def __init__(self):
        """Initialize the reporter."""
        self.logger = logging.getLogger(f"{__name__}.{self.__class__.__name__}")
    
    def generate_report(
        self, 
        results: List[Dict[str, Any]], 
        output_file: Union[str, Path],
        include_summary: bool = True,
        format_override: Optional[str] = None
    ) -> None:
        """
        Generate a report from file analysis results.
        
        Args:
            results: List of analysis results, one per file
            output_file: Path to output file
            include_summary: Whether to include a summary sheet/section
            format_override: Override the output format detection
                            ('excel', 'csv', 'html', 'json')
        """
        # Convert path to string if needed
        output_path = str(output_file) if isinstance(output_file, Path) else output_file
        
        # Detect output format from file extension if not overridden
        if format_override:
            output_format = format_override.lower()
        else:
            _, ext = os.path.splitext(output_path)
            ext = ext.lower()
            
            if ext == '.xlsx':
                output_format = 'excel'
            elif ext == '.csv':
                output_format = 'csv'
            elif ext == '.html' or ext == '.htm':
                output_format = 'html'
            elif ext == '.json':
                output_format = 'json'
            else:
                # Default to Excel
                output_format = 'excel'
                if ext != '.xlsx':
                    output_path += '.xlsx'
                    self.logger.warning(f"No format specified, defaulting to Excel: {output_path}")
        
        # Convert results to a DataFrame
        df = self._prepare_dataframe(results)
        
        # Generate the report in the appropriate format
        if output_format == 'excel':
            self._write_excel_report(df, output_path, include_summary, results)
        elif output_format == 'csv':
            self._write_csv_report(df, output_path)
        elif output_format == 'html':
            self._write_html_report(df, output_path, include_summary, results)
        elif output_format == 'json':
            self._write_json_report(results, output_path)
        else:
            raise ValueError(f"Unsupported output format: {output_format}")
        
        self.logger.info(f"Report generated: {output_path}")
    
    def _prepare_dataframe(self, results: List[Dict[str, Any]]) -> pd.DataFrame:
        """
        Prepare a pandas DataFrame from the analysis results.
        
        Args:
            results: List of analysis results
            
        Returns:
            Formatted pandas DataFrame
        """
        # Convert results to a DataFrame
        df = pd.DataFrame(results)
        
        # Clean up the DataFrame
        
        # Convert date columns to datetime if they're not already
        date_columns = ['modified_date', 'created_date', 'accessed_date']
        for col in date_columns:
            if col in df.columns:
                if df[col].dtype == 'object':
                    try:
                        df[col] = pd.to_datetime(df[col], errors='coerce')
                    except Exception as e:
                        self.logger.warning(f"Error converting {col} to datetime: {str(e)}")
        
        # Handle nested dictionaries (flatten them for display)
        # Columns like 'properties', 'sheets_info', 'column_types'
        nested_columns = [
            col for col in df.columns 
            if df[col].dtype == 'object' and df[col].apply(lambda x: isinstance(x, dict)).any()
        ]
        
        for col in nested_columns:
            # Convert to string representation for display
            df[col] = df[col].apply(lambda x: json.dumps(x, default=str) if isinstance(x, dict) else x)
        
        # Handle lists (convert to string representation)
        list_columns = [
            col for col in df.columns 
            if df[col].dtype == 'object' and df[col].apply(lambda x: isinstance(x, list)).any()
        ]
        
        for col in list_columns:
            df[col] = df[col].apply(lambda x: ', '.join(map(str, x)) if isinstance(x, list) else x)
        
        # Organize columns in a logical order
        column_order = []
        
        # Start with columns from the predefined order
        for col in REPORT_COLUMN_ORDER:
            if col in df.columns:
                column_order.append(col)
        
        # Add any remaining columns
        for col in df.columns:
            if col not in column_order:
                column_order.append(col)
        
        # Reorder the DataFrame
        df = df[column_order]
        
        return df
    
    def _write_excel_report(
        self, 
        df: pd.DataFrame, 
        output_file: str,
        include_summary: bool = True,
        raw_results: Optional[List[Dict[str, Any]]] = None
    ) -> None:
        """
        Write the report to an Excel file.
        
        Args:
            df: DataFrame containing the report data
            output_file: Path to the output Excel file
            include_summary: Whether to include a summary sheet
            raw_results: Original analysis results for summary statistics
        """
        # Create Excel writer
        with pd.ExcelWriter(output_file, engine='openpyxl') as writer:
            # Write the main data sheet
            df.to_excel(writer, sheet_name='File Analysis', index=False)
            
            # Format the main sheet if openpyxl is available
            if OPENPYXL_AVAILABLE:
                self._format_excel_sheet(writer.sheets['File Analysis'], df)
            
            # Create summary sheet if requested
            if include_summary and raw_results:
                self._add_summary_sheet(writer, raw_results)
    
    def _format_excel_sheet(self, worksheet, df: pd.DataFrame) -> None:
        """
        Format an Excel worksheet for better readability.
        
        Args:
            worksheet: openpyxl worksheet object
            df: DataFrame containing the data
        """
        # Define styles
        header_font = Font(bold=True)
        header_fill = PatternFill(
            start_color=EXCEL_HEADER_STYLE.get("fill_color", "E9ECEF"),
            end_color=EXCEL_HEADER_STYLE.get("fill_color", "E9ECEF"),
            fill_type='solid'
        )
        border = Border(
            left=Side(style='thin'),
            right=Side(style='thin'),
            top=Side(style='thin'),
            bottom=Side(style='thin')
        )
        
        # Format headers
        for col_num, column in enumerate(df.columns, 1):
            cell = worksheet.cell(row=1, column=col_num)
            cell.font = header_font
            cell.fill = header_fill
            cell.border = border
            
            # Center the header
            cell.alignment = Alignment(horizontal='center')
        
        # Auto-adjust column widths
        for idx, column in enumerate(df.columns):
            column_width = max(
                min(df[column].astype(str).str.len().max(), 50),  # Max width 50
                len(str(column)) + 2
            )
            worksheet.column_dimensions[get_column_letter(idx + 1)].width = column_width
        
        # Format date columns
        date_columns = ['modified_date', 'created_date', 'accessed_date']
        date_format = "yyyy-mm-dd hh:mm:ss"
        
        for col_idx, column in enumerate(df.columns, 1):
            if column in date_columns:
                # Format date cells
                for row_idx in range(2, len(df) + 2):  # Start from row 2 (after header)
                    cell = worksheet.cell(row=row_idx, column=col_idx)
                    if cell.value:
                        cell.number_format = date_format
        
        # Add borders and zebra striping
        for row_idx in range(2, len(df) + 2):  # Start from row 2 (after header)
            # Add light gray background to even rows for zebra striping
            if row_idx % 2 == 0:
                for col_idx in range(1, len(df.columns) + 1):
                    cell = worksheet.cell(row=row_idx, column=col_idx)
                    cell.fill = PatternFill(
                        start_color="F5F5F5",
                        end_color="F5F5F5",
                        fill_type='solid'
                    )
            
            # Add borders to all cells
            for col_idx in range(1, len(df.columns) + 1):
                cell = worksheet.cell(row=row_idx, column=col_idx)
                cell.border = border
        
        # Freeze the header row
        worksheet.freeze_panes = 'A2'
    
    def _add_summary_sheet(
        self, 
        writer: pd.ExcelWriter,
        results: List[Dict[str, Any]]
    ) -> None:
        """
        Add a summary sheet to the Excel report.
        
        Args:
            writer: pandas ExcelWriter object
            results: Original analysis results for summary statistics
        """
        # Calculate summary statistics
        summary_data = []
        
        # Count file types
        total_files = len(results)
        excel_files = sum(1 for r in results if r.get('is_excel', False))
        csv_files = sum(1 for r in results if r.get('extension') == '.csv')
        other_files = total_files - excel_files - csv_files
        
        # Count special features
        vba_files = sum(1 for r in results if r.get('has_vba', False))
        pivot_files = sum(1 for r in results if r.get('has_pivot_table', False))
        
        # Get file size statistics
        if 'size_kb' in results[0]:
            total_size_kb = sum(r.get('size_kb', 0) for r in results)
            avg_size_kb = total_size_kb / total_files if total_files > 0 else 0
            max_size_kb = max(r.get('size_kb', 0) for r in results)
            min_size_kb = min(r.get('size_kb', 0) for r in results)
        else:
            total_size_kb = avg_size_kb = max_size_kb = min_size_kb = 0
        
        # Build summary data
        summary_data.extend([
            {'Category': 'Total Files', 'Count': total_files, 'Percentage': '100%'},
            {'Category': 'Excel Files', 'Count': excel_files, 
             'Percentage': f"{excel_files/total_files*100:.1f}%" if total_files > 0 else "0%"},
            {'Category': 'CSV Files', 'Count': csv_files, 
             'Percentage': f"{csv_files/total_files*100:.1f}%" if total_files > 0 else "0%"},
            {'Category': 'Other Files', 'Count': other_files, 
             'Percentage': f"{other_files/total_files*100:.1f}%" if total_files > 0 else "0%"},
            {'Category': '', 'Count': '', 'Percentage': ''},
            {'Category': 'Files with VBA', 'Count': vba_files, 
             'Percentage': f"{vba_files/total_files*100:.1f}%" if total_files > 0 else "0%"},
            {'Category': 'Files with PivotTables', 'Count': pivot_files, 
             'Percentage': f"{pivot_files/total_files*100:.1f}%" if total_files > 0 else "0%"},
            {'Category': '', 'Count': '', 'Percentage': ''},
            {'Category': 'Total Size', 'Count': f"{total_size_kb:.2f} KB", 'Percentage': ''},
            {'Category': 'Average Size', 'Count': f"{avg_size_kb:.2f} KB", 'Percentage': ''},
            {'Category': 'Maximum Size', 'Count': f"{max_size_kb:.2f} KB", 'Percentage': ''},
            {'Category': 'Minimum Size', 'Count': f"{min_size_kb:.2f} KB", 'Percentage': ''},
        ])
        
        # Add Excel format statistics
        if excel_files > 0:
            xlsx_files = sum(1 for r in results if r.get('extension') == '.xlsx')
            xlsm_files = sum(1 for r in results if r.get('extension') == '.xlsm')
            xls_files = sum(1 for r in results if r.get('extension') == '.xls')
            xlsb_files = sum(1 for r in results if r.get('extension') == '.xlsb')
            
            summary_data.extend([
                {'Category': '', 'Count': '', 'Percentage': ''},
                {'Category': 'Excel Format Breakdown', 'Count': '', 'Percentage': ''},
                {'Category': 'XLSX Files', 'Count': xlsx_files, 
                 'Percentage': f"{xlsx_files/excel_files*100:.1f}%" if excel_files > 0 else "0%"},
                {'Category': 'XLSM Files', 'Count': xlsm_files, 
                 'Percentage': f"{xlsm_files/excel_files*100:.1f}%" if excel_files > 0 else "0%"},
                {'Category': 'XLS Files', 'Count': xls_files, 
                 'Percentage': f"{xls_files/excel_files*100:.1f}%" if excel_files > 0 else "0%"},
                {'Category': 'XLSB Files', 'Count': xlsb_files, 
                 'Percentage': f"{xlsb_files/excel_files*100:.1f}%" if excel_files > 0 else "0%"},
            ])
        
        # Add CSV format statistics
        if csv_files > 0:
            # Count CSV encodings
            encodings = {}
            for r in results:
                if r.get('extension') == '.csv' and 'encoding' in r:
                    enc = r['encoding']
                    encodings[enc] = encodings.get(enc, 0) + 1
            
            # Add to summary
            summary_data.extend([
                {'Category': '', 'Count': '', 'Percentage': ''},
                {'Category': 'CSV Encoding Breakdown', 'Count': '', 'Percentage': ''},
            ])
            
            for enc, count in encodings.items():
                summary_data.append({
                    'Category': f"Encoding: {enc}",
                    'Count': count,
                    'Percentage': f"{count/csv_files*100:.1f}%" if csv_files > 0 else "0%"
                })
            
            # Count CSV delimiters
            delimiters = {}
            for r in results:
                if r.get('extension') == '.csv' and 'delimiter' in r:
                    delimiter = repr(r['delimiter'])  # Use repr to show whitespace chars
                    delimiters[delimiter] = delimiters.get(delimiter, 0) + 1
            
            # Add to summary
            summary_data.extend([
                {'Category': '', 'Count': '', 'Percentage': ''},
                {'Category': 'CSV Delimiter Breakdown', 'Count': '', 'Percentage': ''},
            ])
            
            for delimiter, count in delimiters.items():
                summary_data.append({
                    'Category': f"Delimiter: {delimiter}",
                    'Count': count,
                    'Percentage': f"{count/csv_files*100:.1f}%" if csv_files > 0 else "0%"
                })
        
        # Create summary DataFrame
        summary_df = pd.DataFrame(summary_data)
        
        # Write to Excel
        summary_df.to_excel(writer, sheet_name='Summary', index=False)
        
        # Format the summary sheet if openpyxl is available
        if OPENPYXL_AVAILABLE:
            sheet = writer.sheets['Summary']
            
            # Set column widths
            sheet.column_dimensions['A'].width = 30
            sheet.column_dimensions['B'].width = 15
            sheet.column_dimensions['C'].width = 15
            
            # Format header
            for col_num in range(1, 4):
                cell = sheet.cell(row=1, column=col_num)
                cell.font = Font(bold=True)
                cell.fill = PatternFill(
                    start_color="E9ECEF",
                    end_color="E9ECEF",
                    fill_type='solid'
                )
                cell.border = Border(
                    left=Side(style='thin'),
                    right=Side(style='thin'),
                    top=Side(style='thin'),
                    bottom=Side(style='thin')
                )
            
            # Format section titles
            for row_num in range(2, len(summary_data) + 2):
                cell = sheet.cell(row=row_num, column=1)
                if cell.value and ':' not in str(cell.value) and cell.value != '':
                    # This is a section title
                    cell.font = Font(bold=True)
            
            # Add report generation timestamp
            timestamp_row = len(summary_data) + 3
            sheet.cell(row=timestamp_row, column=1).value = 'Report Generated:'
            sheet.cell(row=timestamp_row, column=2).value = datetime.now().strftime('%Y-%m-%d %H:%M:%S')
    
    def _write_csv_report(self, df: pd.DataFrame, output_file: str) -> None:
        """
        Write the report to a CSV file.
        
        Args:
            df: DataFrame containing the report data
            output_file: Path to the output CSV file
        """
        df.to_csv(output_file, index=False)
    
    def _write_html_report(
        self, 
        df: pd.DataFrame, 
        output_file: str,
        include_summary: bool = True,
        raw_results: Optional[List[Dict[str, Any]]] = None
    ) -> None:
        """
        Write the report to an HTML file.
        
        Args:
            df: DataFrame containing the report data
            output_file: Path to the output HTML file
            include_summary: Whether to include a summary section
            raw_results: Original analysis results for summary statistics
        """
        from report.templates import HTML_REPORT_TEMPLATE
        
        # Generate summary statistics if requested
        summary_html = ""
        if include_summary and raw_results:
            # Count file types
            total_files = len(raw_results)
            excel_files = sum(1 for r in raw_results if r.get('is_excel', False))
            csv_files = sum(1 for r in raw_results if r.get('extension') == '.csv')
            
            # Count special features
            vba_files = sum(1 for r in raw_results if r.get('has_vba', False))
            pivot_files = sum(1 for r in raw_results if r.get('has_pivot_table', False))
            
            # Create summary HTML
            summary_html = f"""
            <div class="summary">
                <h2>Summary</h2>
                <table class="summary-table">
                    <tr><th>Category</th><th>Count</th></tr>
                    <tr><td>Total Files</td><td>{total_files}</td></tr>
                    <tr><td>Excel Files</td><td>{excel_files}</td></tr>
                    <tr><td>CSV Files</td><td>{csv_files}</td></tr>
                    <tr><td>Files with VBA</td><td>{vba_files}</td></tr>
                    <tr><td>Files with PivotTables</td><td>{pivot_files}</td></tr>
                </table>
            </div>
            """
        
        # Convert DataFrame to HTML table
        table_html = df.to_html(index=False, classes="data-table")
        
        # Create header row HTML
        headers_html = ""
        for col in df.columns:
            headers_html += f"<th>{col}</th>\n"
        
        # Create row HTML
        rows_html = ""
        for _, row in df.iterrows():
            row_html = "<tr>"
            for col in df.columns:
                value = row.get(col, "")
                if pd.isna(value):
                    value = ""
                row_html += f"<td>{value}</td>"
            row_html += "</tr>\n"
            rows_html += row_html
        
        # Fill in the template
        html = HTML_REPORT_TEMPLATE.format(
            date=datetime.now().strftime('%Y-%m-%d %H:%M:%S'),
            total_files=len(raw_results) if raw_results else len(df),
            excel_files=sum(1 for r in raw_results if r.get('is_excel', False)) if raw_results else "N/A",
            csv_files=sum(1 for r in raw_results if r.get('extension') == '.csv') if raw_results else "N/A",
            vba_files=sum(1 for r in raw_results if r.get('has_vba', False)) if raw_results else "N/A",
            pivot_files=sum(1 for r in raw_results if r.get('has_pivot_table', False)) if raw_results else "N/A",
            headers=headers_html,
            rows=rows_html,
            summary=summary_html
        )
        
        # Write the HTML file
        with open(output_file, 'w', encoding='utf-8') as f:
            f.write(html)
    
    def _write_json_report(self, results: List[Dict[str, Any]], output_file: str) -> None:
        """
        Write the report to a JSON file.
        
        Args:
            results: List of analysis results
            output_file: Path to the output JSON file
        """
        # Convert datetime objects to strings for JSON serialization
        def json_serializer(obj):
            if isinstance(obj, (datetime, pd.Timestamp)):
                return obj.isoformat()
            raise TypeError(f"Type {type(obj)} not serializable")
        
        # Create report structure
        report = {
            "metadata": {
                "generated_at": datetime.now().isoformat(),
                "file_count": len(results),
                "excel_count": sum(1 for r in results if r.get('is_excel', False)),
                "csv_count": sum(1 for r in results if r.get('extension') == '.csv'),
                "vba_count": sum(1 for r in results if r.get('has_vba', False)),
                "pivot_count": sum(1 for r in results if r.get('has_pivot_table', False)),
            },
            "results": results
        }
        
        # Write the JSON file
        with open(output_file, 'w', encoding='utf-8') as f:
            json.dump(report, f, default=json_serializer, indent=2)